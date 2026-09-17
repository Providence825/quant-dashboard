# -*- coding: utf-8 -*-
import os, sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
from docx import Document
from docx.shared import Inches, Pt
from docx.oxml.ns import qn

BASE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(BASE, "量化交易看板-产品介绍.docx")
DST = os.path.join(BASE, "量化交易看板-产品介绍_new.docx")
SHOTS = os.path.join(BASE, "shots")

doc = Document(SRC)

def set_font(run, name='宋体', size=10.5, bold=False):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name)

def find_para(prefix):
    for p in doc.paragraphs:
        if p.text.strip().startswith(prefix):
            return p
    raise SystemExit('anchor not found: ' + prefix)

def add_caption(anchor, text):
    p = anchor.insert_paragraph_before(style='Normal')
    r = p.add_run(text)
    set_font(r)
    return p

def add_image(anchor, filename, width=5.83):
    p = anchor.insert_paragraph_before(style='Normal')
    p.add_run().add_picture(os.path.join(SHOTS, filename), width=Inches(width))
    return p

# 1b 行业热力图 —— 插在 "2. 自选监控" 之前（即紧跟盯盘大屏图后）
a2 = find_para('2. 自选监控')
add_caption(a2, '1b. 行业热力图 —— 盯盘大屏内的行业板块热力图，红涨绿跌（A股习惯配色），面积越大表示该板块权重越高。')
add_image(a2, '01b-行业热力图.png')

# 2b K线形态 —— 插在 "3. 账户持仓" 之前（紧跟自选监控图后）
a3 = find_para('3. 账户持仓')
add_caption(a3, '2b. K线形态 —— 从自选股打开的K线分析弹窗（以自选第一只 三安光电 600703 为例），主图叠加均线并自动标注常见形态（十字星等），下方为成交量与MACD。')
add_image(a3, '01c-K线形态.png')

# 8 策略信号四大流派细分 —— 插在 "9. 回测风控" 之前（紧跟策略信号图后）
a9 = find_para('9. 回测风控')
add_caption(a9, '策略信号按四大流派细分，各流派实时信号如下：')
fam = [
    ('8a. 短线策略（震荡市）—— 顶部为行情判断栏，含龙头战法、尾盘捡漏、分时回归、逆向做T。', '08b-策略-短线.png'),
    ('8b. 趋势波段策略（趋势市）—— 均线多头突破、通道突破、均线回踩、海龟通道。', '08c-策略-趋势波段.png'),
    ('8c. 来去由心（情绪周期）—— 依市场情绪分买点1/2/3与模块最强梯队。', '08d-策略-来去由心.png'),
    ('8d. 从零大A（强势接力）—— 关键位突破、假突破、延续性龙头、量能强度、科学加仓。', '08e-策略-从零大A.png'),
]
for cap, fn in fam:
    add_caption(a9, cap)
    add_image(a9, fn)

doc.save(DST)

# 验证
d2 = Document(DST)
imgs = len(d2.inline_shapes)
print('paras:', len(d2.paragraphs))
print('inline images:', imgs)
print('expected images: 15  ->', 'OK' if imgs == 15 else 'MISMATCH')
